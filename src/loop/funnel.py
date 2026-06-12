"""Funnel analyzer — localize where each failed criterion's fact died.

For every failed scorer criterion, walk the pipeline funnel:

    NOT_EXTRACTED      fact never became a claim
    NOT_BOUND          claim exists but reached no target
    NOT_IN_PACKET      bound claim dropped by packet selection
    NOT_IN_DELIVERABLE packet carried it, synthesis dropped it
    QUALITY            present in deliverable but failed the criterion

The histogram of death stages across a batch IS the structural diagnosis:
a modal stage is a structural bottleneck; a uniform scatter is noise.

Usage: python -m src.cli funnel <results_dir> [--judge-model MODEL]
"""
from __future__ import annotations

import json
import re
from concurrent.futures import ThreadPoolExecutor, as_completed
from pathlib import Path

STAGES = (
    "NOT_EXTRACTED", "NOT_BOUND", "NOT_IN_PACKET",
    "NOT_IN_DELIVERABLE", "QUALITY",
)

_WORD = re.compile(r"[a-z0-9][a-z0-9.,$%/-]{2,}")
_STOP = frozenset(
    "the and for that with from this not are was were has have been must "
    "should correctly stated required criterion letter agent output fails "
    "failed identifies identified mention mentions explicitly specific "
    "include includes".split()
)


def _terms(text: str) -> set[str]:
    return {w for w in _WORD.findall(text.lower()) if w not in _STOP}


def _overlap(terms: set[str], text: str) -> float:
    if not terms:
        return 0.0
    lower = text.lower()
    return sum(1 for t in terms if t in lower) / len(terms)


def _top_matches(terms: set[str], items: list[tuple[str, str]],
                 k: int) -> list[tuple[str, str]]:
    scored = sorted(items, key=lambda it: -_overlap(terms, it[1]))
    return [it for it in scored[:k] if _overlap(terms, it[1]) > 0]


def _windows(text: str, terms: set[str], k: int = 3,
             width: int = 900) -> list[str]:
    if not text:
        return []
    spans = []
    for i in range(0, len(text), width // 2):
        w = text[i:i + width]
        spans.append((_overlap(terms, w), w))
    spans.sort(key=lambda s: -s[0])
    return [w for score, w in spans[:k] if score > 0]


def analyze_task(task_dir: Path, caller) -> dict | None:
    """Funnel-analyze every failed criterion of one task."""
    scores_path = task_dir / "scores.json"
    loop_dir = task_dir / "loop"
    if not scores_path.exists() or not loop_dir.exists():
        return None
    scores = json.loads(scores_path.read_text(encoding="utf-8"))
    fails = [
        c for c in scores.get("criteria_results", [])
        if c.get("verdict") != "pass"
    ]
    if not fails:
        return {"task": scores.get("task", str(task_dir)), "results": []}

    # Load funnel artifacts
    finals = sorted(loop_dir.glob("board_iter_*final*.json"))
    if not finals:
        return None
    board = json.loads(finals[-1].read_text(encoding="utf-8"))
    claims = [
        (c["id"],
         f"[{c['kind']}] {c['content']} | {c.get('evidence') or ''}",
         bool(c.get("target_refs")))
        for c in board.get("claims", [])
    ]
    packets_text = []
    for p in loop_dir.glob("packets_*.json"):
        packets_text.append(p.read_text(encoding="utf-8"))
    packets_blob = "\n".join(packets_text)
    deliverable = ""
    out_dir = task_dir / "output"
    if out_dir.exists():
        for f in out_dir.iterdir():
            deliverable += _file_text(f) + "\n"

    def classify(crit: dict) -> dict:
        probe = f"{crit.get('title', '')} {crit.get('reasoning', '')}"
        terms = _terms(probe)
        cand_claims = _top_matches(
            terms, [(cid, txt) for cid, txt, _ in claims], 25,
        )
        bound_map = {cid: b for cid, _, b in claims}
        packet_hits = _windows(packets_blob, terms)
        deliv_hits = _windows(deliverable, terms)

        prompt = f"""You are localizing WHERE in a document-analysis pipeline a required fact died. The pipeline: sources are read into claims -> claims are bound to targets -> bound claims enter synthesis packets -> synthesis writes the deliverable.

FAILED CRITERION: {crit.get('title', '')}
SCORER REASONING: {crit.get('reasoning', '')[:400]}

BEST-MATCHING CLAIMS ON THE BOARD (id | content | bound):
{chr(10).join(f'{cid} | {txt[:200]} | bound={bound_map.get(cid)}' for cid, txt in cand_claims) or '(no matching claims)'}

BEST-MATCHING PACKET EXCERPTS (what synthesis saw):
{chr(10).join(h[:400] for h in packet_hits) or '(no matching packet content)'}

BEST-MATCHING DELIVERABLE EXCERPTS:
{chr(10).join(h[:400] for h in deliv_hits) or '(no matching deliverable content)'}

Classify the death stage:
- NOT_EXTRACTED: the needed fact appears in no claim
- NOT_BOUND: claim(s) exist with the fact but bound=False
- NOT_IN_PACKET: bound claim(s) exist but no packet excerpt carries the fact
- NOT_IN_DELIVERABLE: packets carry the fact but the deliverable does not
- QUALITY: the deliverable contains the fact/topic but failed anyway (wrong form, insufficient specificity, scorer judgment)

Return JSON: {{"stage": "<one of the five>", "note": "<one sentence>"}}"""

        result = caller.complete(prompt, max_tokens=512, temperature=0.0,
                                 json_mode=True)
        try:
            parsed = json.loads(result.text)
        except (json.JSONDecodeError, AttributeError):
            parsed = {}
        stage = parsed.get("stage", "")
        if stage not in STAGES:
            stage = "QUALITY"
        return {
            "criterion": crit.get("title", "")[:140],
            "stage": stage,
            "note": str(parsed.get("note", ""))[:200],
        }

    results = []
    with ThreadPoolExecutor(max_workers=10) as pool:
        futures = [pool.submit(classify, c) for c in fails]
        for fut in as_completed(futures):
            try:
                results.append(fut.result())
            except Exception as e:
                results.append({"criterion": "(error)", "stage": "QUALITY",
                                "note": str(e)[:120]})

    out = {"task": scores.get("task", str(task_dir)), "results": results}
    (loop_dir / "funnel_analysis.json").write_text(
        json.dumps(out, indent=1), encoding="utf-8",
    )
    return out


def analyze_batch(results_dir: str, judge_model: str = "gemini-3.1-flash-lite") -> None:
    from ..providers.gemini import GeminiCaller
    caller = GeminiCaller(model=judge_model)
    root = Path(results_dir)
    task_dirs = sorted({p.parent for p in root.rglob("scores.json")})
    print(f"Funnel-analyzing {len(task_dirs)} tasks")

    histogram: dict[str, int] = {s: 0 for s in STAGES}
    per_task = []
    for td in task_dirs:
        out = analyze_task(td, caller)
        if out is None:
            continue
        counts: dict[str, int] = {}
        for r in out["results"]:
            counts[r["stage"]] = counts.get(r["stage"], 0) + 1
            histogram[r["stage"]] += 1
        per_task.append((out["task"], counts))
        print(f"  {out['task'][:64]}: {counts}")

    total = sum(histogram.values())
    print()
    print("=== DEATH-STAGE HISTOGRAM (the structural diagnosis) ===")
    for stage in STAGES:
        n = histogram[stage]
        pct = n / max(total, 1) * 100
        print(f"  {stage:<20} {n:>4}  {pct:>5.1f}%  {'#' * int(pct / 2)}")
    (root / "funnel_report.json").write_text(
        json.dumps({"histogram": histogram,
                    "per_task": [{"task": t, "counts": c} for t, c in per_task]},
                   indent=1),
        encoding="utf-8",
    )
    print(f"\nReport: {root / 'funnel_report.json'}")


def _file_text(path: Path) -> str:
    suffix = path.suffix.lower()
    try:
        if suffix == ".docx":
            from docx import Document as DocxDocument
            return "\n".join(p.text for p in DocxDocument(str(path)).paragraphs)
        if suffix == ".xlsx":
            import openpyxl
            wb = openpyxl.load_workbook(str(path), read_only=True)
            parts = []
            for ws in wb.worksheets:
                for row in ws.iter_rows(values_only=True):
                    parts.append(" | ".join(str(v) for v in row if v is not None))
            return "\n".join(parts)
        return path.read_text(encoding="utf-8", errors="replace")
    except Exception:
        return ""
